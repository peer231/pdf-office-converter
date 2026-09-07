import os
import tempfile
from io import BytesIO

import pymupdf
import pytesseract
import pandas as pd

from PIL import Image
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pdf2docx import Converter


import shutil

# Tesseract configuration
# Windows par local development ke liye
if os.name == "nt":
    pytesseract.pytesseract.tesseract_cmd = (
        r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    )

# Streamlit Cloud / Linux ke liye
else:
    tesseract_path = shutil.which("tesseract")

    if tesseract_path:
        pytesseract.pytesseract.tesseract_cmd = tesseract_path


# ---------------------------------------------------
# CHECK PDF TYPE
# ---------------------------------------------------

def is_scanned_pdf(pdf_bytes):
    """
    Check karta hai PDF normal text PDF hai
    ya scanned/image PDF.
    """

    pdf = pymupdf.open(
        stream=pdf_bytes,
        filetype="pdf"
    )

    total_text = ""

    for page in pdf:

        total_text += page.get_text("text").strip()

    pdf.close()

    # Agar proper text bohat kam mila
    # to scanned PDF consider karenge
    return len(total_text) < 100


# ---------------------------------------------------
# NORMAL PDF -> WORD
# ---------------------------------------------------

def normal_pdf_to_word(pdf_bytes):
    """
    Digital PDF ko pdf2docx ke through Word mein
    convert karta hai.

    Layout, tables, paragraphs etc. preserve
    karne ki koshish karta hai.
    """

    with tempfile.TemporaryDirectory() as temp_dir:

        pdf_path = os.path.join(
            temp_dir,
            "input.pdf"
        )

        docx_path = os.path.join(
            temp_dir,
            "output.docx"
        )

        # PDF save
        with open(pdf_path, "wb") as file:

            file.write(pdf_bytes)

        # Conversion
        converter = Converter(pdf_path)

        converter.convert(
            docx_path,
            start=0,
            end=None
        )

        converter.close()

        # Word read
        with open(docx_path, "rb") as file:

            word_bytes = BytesIO(
                file.read()
            )

        word_bytes.seek(0)

        return word_bytes


# ---------------------------------------------------
# OCR PAGE
# ---------------------------------------------------

def get_ocr_data(page):
    """
    PDF page ko high-quality image mein convert karta hai
    aur OCR word positions return karta hai.
    """

    zoom = 3

    matrix = pymupdf.Matrix(
        zoom,
        zoom
    )

    pix = page.get_pixmap(
        matrix=matrix,
        alpha=False
    )

    image = Image.open(
        BytesIO(
            pix.tobytes("png")
        )
    )

    ocr_data = pytesseract.image_to_data(
        image,
        output_type=pytesseract.Output.DATAFRAME,
        config="--psm 6"
    )

    ocr_data = ocr_data.dropna()

    ocr_data = ocr_data[
        ocr_data["text"].str.strip() != ""
    ]

    return image, ocr_data


# ---------------------------------------------------
# SCANNED PDF -> EDITABLE WORD
# ---------------------------------------------------

def scanned_pdf_to_word(pdf_bytes):
    """
    Scanned PDF ko OCR use karke editable Word
    mein convert karta hai.
    """

    pdf = pymupdf.open(
        stream=pdf_bytes,
        filetype="pdf"
    )

    document = Document()

    # Word margins
    section = document.sections[0]

    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    for page_number, page in enumerate(pdf):

        image, ocr_data = get_ocr_data(page)

        # OCR lines reconstruct karna
        grouped = ocr_data.groupby(
            [
                "block_num",
                "par_num",
                "line_num"
            ]
        )

        lines = []

        for _, line_data in grouped:

            words = line_data.sort_values(
                "left"
            )

            text = " ".join(
                words["text"].astype(str)
            )

            top = int(
                words["top"].min()
            )

            left = int(
                words["left"].min()
            )

            height = int(
                words["height"].max()
            )

            lines.append(
                {
                    "text": text,
                    "top": top,
                    "left": left,
                    "height": height
                }
            )

        # vertical order
        lines.sort(
            key=lambda x: x["top"]
        )

        previous_top = 0

        image_width = image.width

        for line in lines:

            text = line["text"].strip()

            if not text:
                continue

            paragraph = document.add_paragraph()

            # Horizontal position approximate
            left_ratio = (
                line["left"]
                / image_width
            )

            paragraph.paragraph_format.left_indent = Inches(
                left_ratio * 6.5
            )

            # Vertical gaps
            gap = line["top"] - previous_top

            if gap > 60:

                paragraph.paragraph_format.space_before = Pt(
                    min(
                        gap / 5,
                        30
                    )
                )

            run = paragraph.add_run(
                text
            )

            # Font size estimate
            font_size = max(
                8,
                min(
                    16,
                    line["height"] / 3
                )
            )

            run.font.size = Pt(
                font_size
            )

            # Headings detect
            upper_text = text.upper()

            if (
                "AFFIDAVIT" in upper_text
                or text == "Annex-C"
            ):

                run.bold = True

                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

                run.font.size = Pt(15)

            previous_top = line["top"]

        # new page
        if page_number < len(pdf) - 1:

            document.add_page_break()

    pdf.close()

    output = BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output


# ---------------------------------------------------
# MAIN WORD CONVERTER
# ---------------------------------------------------

def create_word(pdf_bytes):
    """
    Automatically PDF type detect karta hai.

    Digital PDF:
        pdf2docx

    Scanned PDF:
        OCR editable Word
    """

    if is_scanned_pdf(pdf_bytes):

        return scanned_pdf_to_word(
            pdf_bytes
        )

    return normal_pdf_to_word(
        pdf_bytes
    )


# ---------------------------------------------------
# EXTRACT TEXT
# ---------------------------------------------------

def extract_text_from_pdf(pdf_bytes):

    pdf = pymupdf.open(
        stream=pdf_bytes,
        filetype="pdf"
    )

    pages = []

    for page_number, page in enumerate(
        pdf,
        start=1
    ):

        text = page.get_text(
            "text"
        ).strip()

        # OCR fallback
        if len(text) < 100:

            zoom = 3

            pix = page.get_pixmap(
                matrix=pymupdf.Matrix(
                    zoom,
                    zoom
                ),
                alpha=False
            )

            image = Image.open(
                BytesIO(
                    pix.tobytes("png")
                )
            )

            text = pytesseract.image_to_string(
                image,
                config="--psm 6"
            )

        pages.append(
            {
                "page": page_number,
                "text": text.strip()
            }
        )

    pdf.close()

    return pages


# ---------------------------------------------------
# CSV
# ---------------------------------------------------

def create_csv(pages):

    rows = []

    for page in pages:

        text = page["text"]

        lines = text.splitlines()

        for line_number, line in enumerate(
            lines,
            start=1
        ):

            if line.strip():

                rows.append(
                    {
                        "page": page["page"],
                        "line": line_number,
                        "text": line.strip()
                    }
                )

    dataframe = pd.DataFrame(
        rows
    )

    output = BytesIO()

    dataframe.to_csv(
        output,
        index=False,
        encoding="utf-8-sig"
    )

    output.seek(0)

    return output
